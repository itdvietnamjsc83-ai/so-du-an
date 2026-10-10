"""Tạo các file Word mẫu có sẵn (src/templates/*.docx) và file khung (khung.docx).
Chạy lại khi muốn sửa mẫu gốc:  python3 scripts/make_templates.py"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

OUT = os.path.join(os.path.dirname(__file__), '..', 'src', 'templates')

def base():
    d = Document()
    s = d.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21)
    s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(2), Cm(2), Cm(3), Cm(1.5)
    st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(13)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(3); st.paragraph_format.line_spacing = 1.15
    return d

def para(d, text='', bold=False, italic=False, align=None, size=None, after=None):
    p = d.add_paragraph()
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
    if align == 'c': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if align == 'r': p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if align == 'j': p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def cell_lines(cell, lines):
    cell.paragraphs[0].text = ''
    for i, (t, b, it) in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t); r.bold = b; r.italic = it; r.font.size = Pt(12.5)

def header(d, so_prefix):
    t = d.add_table(rows=1, cols=2); t.autofit = False
    t.columns[0].width = Cm(6.5); t.columns[1].width = Cm(10)
    for c, w in zip(t.rows[0].cells, (Cm(6.5), Cm(10))): c.width = w
    cell_lines(t.rows[0].cells[0], [('{TuVanQLDA.Ten}', True, False), ('Số: {SoVanBan}', False, False)])
    cell_lines(t.rows[0].cells[1], [('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', True, False), ('Độc lập - Tự do - Hạnh phúc', True, False), ('{DiaDiem}, {Ngay}', False, True)])
    para(d)

def info(d, rows):
    for k, v in rows:
        p = para(d, align='j'); p.add_run(k + ': ').bold = False; p.add_run(v)

def sign(d, left, right):
    t = d.add_table(rows=1, cols=2)
    for c, txt in zip(t.rows[0].cells, (left, right)):
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].add_run(txt).bold = True
        for _ in range(4): c.add_paragraph()

DUAN = [('Dự án', '{DuAn.Ten}'), ('Địa điểm', '{DuAn.DiaDiem}'), ('Chủ đầu tư', '{ChuDauTu.Ten}')]
GOI = [('Gói thầu', '{GoiThau.Ma} - {GoiThau.Ten}'), ('Hợp đồng số', '{HopDong.So} ngày {HopDong.Ngay}'), ('Nhà thầu', '{NhaThau.Ten}')]

def make(name, title, nhom, body):
    d = base(); header(d, '')
    para(d, title.upper(), bold=True, align='c', size=14, after=6)
    body(d)
    d.save(os.path.join(OUT, name + '.docx'))

def gm(d):
    para(d, 'Kính gửi:', bold=True); para(d, '{ThanhPhan}')
    para(d, '{TuVanQLDA.Ten} trân trọng kính mời đại diện các đơn vị tham dự cuộc họp:', align='j')
    info(d, DUAN + [('Thời gian', '{ThoiGianHop}'), ('Địa điểm họp', '{DiaDiemHop}'), ('Nội dung', '{NoiDungHop}')])
    para(d, 'Đề nghị các đơn vị chuẩn bị tài liệu liên quan và cử người tham dự đúng thành phần. Rất mong được sự quan tâm, tham dự đầy đủ, đúng giờ.', align='j')
    sign(d, 'Nơi nhận:\n- Như trên;\n- Lưu: VT.', '{TuVanQLDA.DaiDien}')

def bbhop(d):
    info(d, DUAN + [('Thời gian', '{ThoiGianHop}'), ('Địa điểm', '{DiaDiemHop}'), ('Chủ trì', '{ChuTri}')])
    para(d, 'I. THÀNH PHẦN THAM DỰ', bold=True); para(d, '{ThanhPhan}')
    para(d, 'II. NỘI DUNG CUỘC HỌP', bold=True); para(d, '{NoiDungHop}')
    para(d, 'III. Ý KIẾN CÁC BÊN', bold=True); para(d, '{YKien}')
    para(d, 'IV. KẾT LUẬN VÀ PHÂN CÔNG', bold=True); para(d, '{KetLuan}')
    para(d, 'Cuộc họp kết thúc lúc ...... cùng ngày. Biên bản được các bên thống nhất thông qua.', align='j')
    sign(d, 'THƯ KÝ', 'CHỦ TRÌ')

def bblv(d):
    info(d, DUAN + GOI + [('Địa điểm', '{DiaDiemLamViec}')])
    para(d, 'I. THÀNH PHẦN', bold=True); para(d, '{ThanhPhan}')
    para(d, 'II. NỘI DUNG LÀM VIỆC', bold=True); para(d, '{NoiDungLamViec}')
    para(d, 'III. KẾT QUẢ, THỐNG NHẤT', bold=True); para(d, '{KetLuan}')
    sign(d, 'ĐẠI DIỆN NHÀ THẦU', 'ĐẠI DIỆN {TuVanQLDA.Ten}')

def bbht(d):
    info(d, DUAN + GOI)
    para(d, 'I. THÀNH PHẦN KIỂM TRA', bold=True); para(d, '{ThanhPhan}')
    para(d, 'II. HIỆN TRẠNG', bold=True); para(d, '{HienTrang}')
    para(d, 'III. TỒN TẠI VÀ YÊU CẦU KHẮC PHỤC', bold=True); para(d, '{YeuCau}')
    para(d, 'Thời hạn khắc phục: {ThoiHanKhacPhuc}')
    sign(d, 'ĐẠI DIỆN NHÀ THẦU', 'ĐẠI DIỆN {TuVanQLDA.Ten}')

def nt_cv(d):
    info(d, DUAN + GOI + [('Công việc nghiệm thu', '{CongViecNghiemThu}'), ('Vị trí', '{ViTri}')])
    para(d, 'I. THÀNH PHẦN TRỰC TIẾP NGHIỆM THU', bold=True); para(d, '{ThanhPhan}')
    para(d, 'II. THỜI GIAN NGHIỆM THU', bold=True); para(d, 'Bắt đầu: ............ Kết thúc: ............')
    para(d, 'III. ĐÁNH GIÁ CÔNG VIỆC ĐÃ THỰC HIỆN', bold=True)
    para(d, '1. Căn cứ nghiệm thu: {CanCu}'); para(d, '2. Chất lượng công việc: {DanhGia}'); para(d, '3. Các ý kiến khác: ............')
    para(d, 'IV. KẾT LUẬN', bold=True); para(d, '{KetLuan}')
    sign(d, 'ĐẠI DIỆN NHÀ THẦU THI CÔNG', 'ĐẠI DIỆN GIÁM SÁT / {TuVanQLDA.Ten}')

def nt_gd(d):
    info(d, DUAN + GOI + [('Giai đoạn / bộ phận nghiệm thu', '{BoPhan}')])
    para(d, 'I. THÀNH PHẦN', bold=True); para(d, '{ThanhPhan}')
    para(d, 'II. ĐÁNH GIÁ', bold=True); para(d, '{DanhGia}')
    para(d, 'III. KẾT LUẬN', bold=True); para(d, '{KetLuan}')
    sign(d, 'ĐẠI DIỆN NHÀ THẦU', 'ĐẠI DIỆN CHỦ ĐẦU TƯ / {TuVanQLDA.Ten}')

def nt_kl(d):
    info(d, DUAN + GOI + [('Giá trị hợp đồng', '{HopDong.GiaTri}')])
    para(d, 'I. THÀNH PHẦN', bold=True); para(d, '{ThanhPhan}')
    para(d, 'II. XÁC NHẬN KHỐI LƯỢNG, GIÁ TRỊ ĐỢT {DotNghiemThu.So}', bold=True)
    info(d, [('Giá trị nghiệm thu đợt này', '{DotNghiemThu.GiaTri}'), ('Khấu trừ tạm ứng', '{DotNghiemThu.KhauTruTamUng}'), ('Giữ lại bảo hành', '{DotNghiemThu.GiuBaoHanh}'), ('Giá trị đề nghị thanh toán', '{DotNghiemThu.DeNghiThanhToan}')])
    para(d, '(Chi tiết khối lượng theo bảng tính kèm theo)', italic=True)
    sign(d, 'ĐẠI DIỆN NHÀ THẦU', 'ĐẠI DIỆN CHỦ ĐẦU TƯ / {TuVanQLDA.Ten}')

def nt_ht(d):
    info(d, DUAN + GOI + [('Hạng mục / công trình', '{HangMuc}')])
    para(d, 'I. THÀNH PHẦN', bold=True); para(d, '{ThanhPhan}')
    para(d, 'II. ĐÁNH GIÁ HẠNG MỤC, CÔNG TRÌNH', bold=True); para(d, '{DanhGia}')
    para(d, 'III. KẾT LUẬN', bold=True); para(d, 'Chấp thuận nghiệm thu hoàn thành để đưa vào sử dụng. {KetLuan}')
    sign(d, 'ĐẠI DIỆN NHÀ THẦU', 'ĐẠI DIỆN CHỦ ĐẦU TƯ')

def bg_mb(d):
    info(d, DUAN + GOI)
    para(d, 'I. THÀNH PHẦN', bold=True); para(d, '{ThanhPhan}')
    para(d, 'II. NỘI DUNG BÀN GIAO', bold=True); para(d, '{NoiDungBanGiao}')
    para(d, 'Kể từ ngày ký biên bản, nhà thầu chịu trách nhiệm quản lý mặt bằng được bàn giao.', align='j')
    sign(d, 'BÊN NHẬN', 'BÊN GIAO')

def bg_ct(d):
    info(d, DUAN + GOI)
    para(d, 'I. THÀNH PHẦN', bold=True); para(d, '{ThanhPhan}')
    para(d, 'II. NỘI DUNG BÀN GIAO', bold=True); para(d, '{NoiDungBanGiao}')
    para(d, 'III. HỒ SƠ KÈM THEO', bold=True); para(d, '{HoSoKemTheo}')
    sign(d, 'BÊN NHẬN', 'BÊN GIAO')

def cv_dd(d):
    para(d, 'V/v: Đôn đốc tiến độ thi công {GoiThau.Ma}', italic=True)
    para(d, 'Kính gửi: {NhaThau.Ten}', bold=True, align='c')
    para(d, 'Căn cứ Hợp đồng số {HopDong.So} ngày {HopDong.Ngay} thực hiện gói thầu {GoiThau.Ten} thuộc dự án {DuAn.Ten};', align='j')
    para(d, 'Qua theo dõi, các công việc sau đang chậm so với tiến độ đã cam kết:', align='j')
    para(d, '{ViecQuaHan}')
    para(d, '{TuVanQLDA.Ten} đề nghị {NhaThau.Ten} khẩn trương bổ sung nhân lực, thiết bị, có biện pháp khắc phục và báo cáo kế hoạch hoàn thành trước ngày {HanBaoCao}.', align='j')
    sign(d, 'Nơi nhận:\n- Như trên;\n- Chủ đầu tư (b/c);\n- Lưu: VT.', '{TuVanQLDA.DaiDien}')

def tt(d):
    para(d, 'Kính gửi: {ChuDauTu.Ten}', bold=True, align='c')
    info(d, DUAN)
    para(d, 'Nội dung trình: {NoiDungTrinh}', align='j')
    para(d, 'Căn cứ: {CanCu}', align='j')
    para(d, 'Đề xuất: {DeXuat}', align='j')
    para(d, 'Kính trình {ChuDauTu.Ten} xem xét, phê duyệt.', align='j')
    sign(d, 'Nơi nhận:\n- Như trên;\n- Lưu: VT.', '{TuVanQLDA.DaiDien}')

TPL = [
    ('BM01', 'Giấy mời họp', 'Họp', gm), ('BM02', 'Biên bản họp', 'Họp', bbhop),
    ('BM03', 'Biên bản làm việc', 'Làm việc', bblv), ('BM04', 'Biên bản kiểm tra hiện trường', 'Làm việc', bbht),
    ('BM05', 'Biên bản nghiệm thu công việc xây dựng', 'Nghiệm thu', nt_cv), ('BM06', 'Biên bản nghiệm thu giai đoạn, bộ phận', 'Nghiệm thu', nt_gd),
    ('BM07', 'Biên bản nghiệm thu khối lượng để thanh toán', 'Nghiệm thu', nt_kl), ('BM08', 'Biên bản nghiệm thu hoàn thành đưa vào sử dụng', 'Nghiệm thu', nt_ht),
    ('BM09', 'Biên bản bàn giao mặt bằng', 'Bàn giao', bg_mb), ('BM10', 'Biên bản bàn giao công trình', 'Bàn giao', bg_ct),
    ('BM11', 'Công văn đôn đốc tiến độ', 'Văn bản', cv_dd), ('BM12', 'Tờ trình', 'Văn bản', tt),
]
for i, t, n, f in TPL: make(i, t, n, f)

# Khung để tạo mẫu mới: [[Tieu]] là tiêu đề, [[@body]] là phần thân do phần mềm sinh.
d = base(); header(d, '')
para(d, '[[Tieu]]', bold=True, align='c', size=14, after=6)
para(d, '[[@body]]')
sign(d, 'ĐẠI DIỆN CÁC BÊN', '{TuVanQLDA.DaiDien}')
d.save(os.path.join(OUT, 'khung.docx'))
print('ok', len(TPL) + 1)
